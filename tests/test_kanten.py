"""Randfälle: Original schützen, Bezüge nach dem Löschen von Zeilen, Diagrammblätter."""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
from openpyxl.worksheet.table import Table

from conftest import FIX, sha
from excel_skills import laden


def _basis(ws) -> None:
    """Kopfzeile in Zeile 4, Leerzeile 7. Bereinigen löscht Zeilen 1-3 und 7, C8 wird zu C4."""
    for zeile in (["Titel"], [None], [None], ["Kunde", "Ort", "Betrag"], ["A", "Kiel", 10], ["B", "Ulm", 20], [None], ["C", "Hof", 30]):
        ws.append(zeile)


def _mappe(titel: str = "Daten"):
    wb = Workbook()
    ws = wb.active
    ws.title = titel
    _basis(ws)
    return wb, ws


@pytest.mark.parametrize("befehl,endung", [
    ("profil", ".md"), ("formeln", ".md"), ("csv", ".csv"),
])
def test_original_bleibt_byte_gleich_bei_berichten(lauf, tmp_path, befehl, endung):
    quelle = FIX / "umsatz_bericht.xlsx"
    vorher = sha(quelle)
    code, _ = lauf(befehl, quelle, "--out", tmp_path / f"aus{endung}")
    assert code == 0
    assert sha(quelle) == vorher


def test_original_bleibt_byte_gleich_bei_vergleich_und_stapeln(lauf, tmp_path):
    alt, neu = FIX / "kunden_liste.xlsx", FIX / "kunden_liste_neu.xlsx"
    monate = sorted((FIX / "monate").iterdir())
    vorher = [sha(p) for p in (alt, neu, *monate)]
    assert lauf("vergleich", alt, neu, "--schluessel", "Kundennummer", "--out", tmp_path / "v.md")[0] == 0
    assert lauf("stapeln", FIX / "monate", "--out", tmp_path / "g.xlsx")[0] == 0
    assert [sha(p) for p in (alt, neu, *monate)] == vorher


def test_out_per_hardlink_oder_andere_schreibweise_wird_abgelehnt(lauf, tmp_path):
    quelle = tmp_path / "k.xlsx"
    shutil.copy(FIX / "kunden_liste.xlsx", quelle)
    vorher = sha(quelle)
    (tmp_path / "link.xlsx").hardlink_to(quelle)
    ziele = [tmp_path / "link.xlsx"]
    if (tmp_path / "K.xlsx").exists():  # Dateisystem ohne Unterscheidung von Groß- und Kleinschreibung
        ziele.append(tmp_path / "K.xlsx")
    for ziel in ziele:
        code, erg = lauf("bereinigen", quelle, "--out", ziel)
        assert code == 2, ziel
        assert "Eingabedatei" in erg["fehler"]
    assert sha(quelle) == vorher


def test_diagrammblatt_per_blatt_ist_sauberer_fehler(lauf, tmp_path):
    wb, ws = _mappe()
    cs = wb.create_chartsheet("Diagramm")
    ch = BarChart()
    ch.add_data(Reference(ws, min_col=3, min_row=4, max_row=8))
    cs.add_chart(ch)
    wb.save(tmp_path / "chart.xlsx")
    for befehl, endung in (("csv", ".csv"), ("bereinigen", ".xlsx")):
        code, erg = lauf(befehl, tmp_path / "chart.xlsx", "--blatt", "Diagramm", "--out", tmp_path / f"aus{endung}")
        assert code == 2
        assert "Tabellenblatt 'Diagramm' fehlt" in erg["fehler"]


def test_bereinigen_verschiebt_matrixformel_tabelle_link_namen_druckbereich(lauf, tmp_path):
    wb, ws = _mappe()
    wb.create_sheet("N")
    ws["E8"] = ArrayFormula("E8", "=SUM(N!A1:A2*N!B1:B2)")
    ws["B8"].hyperlink = "https://example.com/hof"
    ws.add_table(Table(displayName="T1", ref="A4:C8"))
    ws.defined_names["Lokal"] = DefinedName("Lokal", attr_text="Daten!$C$8")
    wb.defined_names["Global"] = DefinedName("Global", attr_text="Daten!$C$8")
    ws.print_area = "A4:C8"
    wb.save(tmp_path / "m.xlsx")

    code, erg = lauf("bereinigen", tmp_path / "m.xlsx", "--out", tmp_path / "m_out.xlsx")
    assert code == 0, erg
    wb2 = load_workbook(tmp_path / "m_out.xlsx")
    ws2 = wb2["Daten"]
    assert ws2["E4"].value.ref == "E4"
    assert ws2["B4"].value == "Hof" and ws2["B4"].hyperlink.ref == "B4"
    assert ws2.tables["T1"].ref == "A1:C4"
    assert wb2.defined_names["Global"].attr_text == "Daten!$C$4"
    assert ws2.defined_names["Lokal"].attr_text == "Daten!$C$4"
    assert ws2.print_area.endswith("$A$1:$C$4")


def test_bereinigen_apostroph_im_blattnamen(lauf, tmp_path):
    wb, _ = _mappe(titel="Kunde's Liste")
    wb.create_sheet("N")["A1"] = "='Kunde''s Liste'!C8"
    wb.save(tmp_path / "a.xlsx")
    assert lauf("bereinigen", tmp_path / "a.xlsx", "--out", tmp_path / "a_out.xlsx")[0] == 0
    assert load_workbook(tmp_path / "a_out.xlsx")["N"]["A1"].value == "='Kunde''s Liste'!C4"


def test_stapeln_lehnt_formeln_bei_anderer_spaltenreihenfolge_ab(lauf, tmp_path):
    ordner = tmp_path / "monate"
    ordner.mkdir()
    for name, kopf, zeile in (("a.xlsx", ["Artikel", "Menge", "Preis", "Summe"], ["X", 2, 3, "=B2*C2"]),
                              ("b.xlsx", ["Menge", "Artikel", "Preis", "Summe"], [4, "Y", 5, "=A2*C2"])):
        wb = Workbook()
        wb.active.append(kopf)
        wb.active.append(zeile)
        wb.save(ordner / name)
    code, erg = lauf("stapeln", ordner, "--out", tmp_path / "g.xlsx")
    assert code == 3
    assert erg["details"][0]["datei"] == "b.xlsx"
    assert not (tmp_path / "g.xlsx").exists()


def test_formeltext_einer_datentabelle():
    zelle = Workbook().active["A1"]
    zelle.value = DataTableFormula(ref="A1:B2", dt2D=True, r1="C1", r2="C2")
    assert laden.formeltext(zelle) == "<DataTableFormula>"


def test_firmenname_mit_summenwort_bleibt(lauf, tmp_path):
    wb = Workbook()
    ws = wb.active
    for zeile in (["Kunde", "Ort", "Betrag"], ["Gesamtschule Nord", "Kiel", 10], ["Totalbau GmbH", "Ulm", 20],
                  ["Summerfield Ltd", "Hof", 5], ["Gesamtsumme", None, 35]):
        ws.append(zeile)
    wb.save(tmp_path / "s.xlsx")
    code, erg = lauf("bereinigen", tmp_path / "s.xlsx", "--out", tmp_path / "s_out.xlsx")
    assert code == 0
    assert erg["zeilen_geloescht"] == 1
    namen = [r[0] for r in load_workbook(tmp_path / "s_out.xlsx").active.iter_rows(values_only=True)]
    assert namen == ["Kunde", "Gesamtschule Nord", "Totalbau GmbH", "Summerfield Ltd"]
