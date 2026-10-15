from __future__ import annotations

import json
import shutil
import zipfile

from openpyxl import Workbook, load_workbook

from conftest import FIX, sha


def test_profil_bericht_und_tabelle(lauf, tmp_path):
    code, erg = lauf("profil", FIX / "umsatz_bericht.xlsx", "--out", tmp_path / "p.md")
    assert code == 0
    assert erg["blaetter"] == {"Umsatz": "bericht", "Notizen": "formular"}
    daten = json.loads((tmp_path / "p.json").read_text())
    umsatz = daten["blaetter"][0]
    assert umsatz["kopfzeile"]["zeile"] == 4
    assert umsatz["summenzeilen"] == [8, 12, 13]
    assert {h["zelle"] for h in umsatz["hinweise"]} == {"E7", "C10"}
    code, erg = lauf("profil", FIX / "kunden_liste.xlsx", "--out", tmp_path / "k.md")
    assert erg["blaetter"] == {"Kunden": "tabelle"}
    plz = [s for s in json.loads((tmp_path / "k.json").read_text())["blaetter"][0]["spalten"] if s["name"] == "PLZ"][0]
    assert plz["typ"] == "text"


def test_bereinigen_original_bleibt_und_formeln_wandern(lauf, tmp_path):
    quelle = FIX / "umsatz_bericht.xlsx"
    vorher = sha(quelle)
    out = tmp_path / "sauber.xlsx"
    code, erg = lauf("bereinigen", quelle, "--out", out)
    assert code == 0
    assert sha(quelle) == vorher
    wb = load_workbook(out)
    ws = wb["Umsatz"]
    zeilen = list(ws.iter_rows(values_only=True))
    assert zeilen[0] == ("Region", "Kunde", "Monat", "Menge", "Umsatz", "Menge doppelt")
    assert len(zeilen) == 6
    assert not ws.merged_cells.ranges
    assert [z[0] for z in zeilen[1:]] == ["Nord", "Nord", "Nord", "Süd", "Süd"]
    assert zeilen[1][1] == "Muster Maschinenbau GmbH"
    assert zeilen[1][4] == 1234.56 and zeilen[5][4] == -250.75
    assert zeilen[3][4] == "1.234"
    assert zeilen[4][2] == "01/31/2026"
    assert [z[5] for z in zeilen[1:]] == ["=D2*2", "=D3*2", "=D4*2", "=D5*2", "=D6*2"]
    assert wb["Notizen"]["B1"].value == "=Umsatz!D2"
    assert wb["Notizen"]["B2"].value == "=Umsatz!D6"
    protokoll = json.loads((tmp_path / "sauber.aenderungen.json").read_text())
    assert [(z["zeile"], z["regel"]) for z in protokoll["geloeschte_zeilen"]] == [
        (1, "kopfzeile"), (2, "kopfzeile"), (3, "kopfzeile"), (8, "summen"), (9, "leerzeilen"), (12, "summen"), (13, "summen")]
    regeln = {a["regel"] for a in protokoll["aenderungen"]}
    assert regeln == {"verbunden", "trimmen", "deutsche-zahl", "datum"}


def test_bereinigen_ohne_summen_passt_bereiche_an(lauf, tmp_path):
    out = tmp_path / "s.xlsx"
    code, _ = lauf("bereinigen", FIX / "umsatz_bericht.xlsx", "--ohne", "summen", "--out", out)
    assert code == 0
    ws = load_workbook(out)["Umsatz"]
    formeln = [c.value for c in ws["D"] if isinstance(c.value, str) and c.value.startswith("=")]
    assert formeln == ["=SUM(D2:D4)", "=SUM(D6:D7)", "=D5+D8"]


def test_bereinigen_bricht_ab_wenn_formel_auf_entfallende_zeile_zeigt(lauf, tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Kunde", "Ort", "Betrag", "Anteil"])
    ws.append(["A GmbH", "Kiel", 10, "=C2/C4"])
    ws.append(["B GmbH", "Ulm", 30, "=C3/C4"])
    ws.append(["Summe", None, "=SUM(C2:C3)", None])
    ws.append(["C GmbH", "Hof", 5, "=C5/C4"])
    quelle = tmp_path / "abh.xlsx"
    wb.save(quelle)
    out = tmp_path / "x.xlsx"
    code, erg = lauf("bereinigen", quelle, "--out", out)
    assert code == 3
    assert not out.exists()
    assert {d["zelle"] for d in erg["details"]} == {"D2", "D3", "D5"}


def test_out_gleich_eingabe_wird_abgelehnt(lauf, tmp_path):
    kopie = tmp_path / "k.xlsx"
    shutil.copy(FIX / "kunden_liste.xlsx", kopie)
    vorher = sha(kopie)
    code, erg = lauf("bereinigen", kopie, "--out", kopie)
    assert code == 2 and "nie überschrieben" in erg["fehler"]
    code, _ = lauf("bereinigen", kopie)
    assert code == 2
    code, _ = lauf("csv", kopie, "--out", kopie)
    assert code == 2
    assert sha(kopie) == vorher


def test_formeln(lauf, tmp_path):
    quelle = FIX / "formeln_kaputt.xlsx"
    vorher = sha(quelle)
    code, erg = lauf("formeln", quelle, "--out", tmp_path / "f.md")
    assert code == 0 and sha(quelle) == vorher
    assert (erg["fehler"], erg["harte_werte"], erg["formel_fehlt"], erg["externe_verknuepfungen"]) == (4, 1, 1, 1)
    daten = json.loads((tmp_path / "f.json").read_text())
    kalk = daten["blaetter"][0]
    nv = [f for f in kalk["fehler"] if f["zelle"] == "E5"][0]
    assert nv["deutsch"] == "#NV" and nv["hinweis"].startswith("SVERWEIS")
    assert kalk["harte_werte"][0]["zelle"] == "D4"
    assert kalk["formel_fehlt"][0]["zelle"] == "D7"
    assert daten["blaetter"][1]["externe_verknuepfungen"][0]["zelle"] == "C2"


def test_vergleich_mit_schluessel(lauf, tmp_path):
    code, erg = lauf("vergleich", FIX / "preise_alt.xlsx", FIX / "preise_neu.xlsx", "--schluessel", "Artikelnummer", "--out", tmp_path / "v.md")
    assert code == 0
    assert (erg["neu"], erg["geloescht"], erg["geaendert"]) == (1, 1, 2)
    daten = json.loads((tmp_path / "v.json").read_text())
    assert daten["neu"][0]["schluessel"] == "4715"
    assert daten["geloescht"][0]["schluessel"] == "4713"
    assert {(g["schluessel"], g["spalte"]) for g in daten["geaendert"]} == {("4711", "Preis"), ("4712", "Bezeichnung")}


def test_vergleich_fehlender_schluessel(lauf, tmp_path):
    code, erg = lauf("vergleich", FIX / "preise_alt.xlsx", FIX / "preise_neu.xlsx", "--schluessel", "SKU", "--out", tmp_path / "v.md")
    assert code == 2 and "SKU" in erg["fehler"]


def test_stapeln(lauf, tmp_path):
    out = tmp_path / "gesamt.xlsx"
    code, erg = lauf("stapeln", FIX / "monate", "--out", out)
    assert code == 0
    assert (erg["dateien"], erg["zeilen_gesamt"], erg["uebersprungen"]) == (3, 5, 1)
    zeilen = list(load_workbook(out).active.iter_rows(values_only=True))
    assert zeilen[0] == ("Datum", "Kunde", "Betrag", "quelle_datei")
    assert [z[3] for z in zeilen[1:]] == ["2026-01.xlsx"] * 2 + ["2026-02.xlsx"] + ["2026-03.xlsx"] * 2


def test_stapeln_bricht_bei_abweichender_spalte_ab(lauf, tmp_path):
    out = tmp_path / "g.xlsx"
    code, erg = lauf("stapeln", FIX / "monate_abweichend", "--out", out)
    assert code == 3 and not out.exists()
    assert erg["details"] == [{"datei": "2026-02.xlsx", "fehlt": ["Betrag"], "zusaetzlich": ["Netto"], "doppelt": []}]


def test_stapeln_verschiebt_formeln_und_lehnt_blattbezug_ab(lauf, tmp_path):
    ordner = tmp_path / "in"
    ordner.mkdir()
    for name, extra in (("a.xlsx", None), ("b.xlsx", None)):
        wb = Workbook()
        ws = wb.active
        ws.append(["Artikel", "Menge", "Preis", "Summe"])
        ws.append(["X", 2, 3, "=B2*C2"])
        ws.append(["Y", 4, 5, "=B3*C3"])
        wb.save(ordner / name)
    code, _ = lauf("stapeln", ordner, "--out", tmp_path / "g.xlsx")
    assert code == 0
    summen = [r[3] for r in load_workbook(tmp_path / "g.xlsx").active.iter_rows(min_row=2, values_only=True)]
    assert summen == ["=B2*C2", "=B3*C3", "=B4*C4", "=B5*C5"]
    wb = load_workbook(ordner / "b.xlsx")
    wb.active["D2"] = "=Preise!B2"
    wb.save(ordner / "b.xlsx")
    code, erg = lauf("stapeln", ordner, "--out", tmp_path / "g2.xlsx")
    assert code == 3 and not (tmp_path / "g2.xlsx").exists()
    code, _ = lauf("stapeln", ordner, "--werte", "--out", tmp_path / "g3.xlsx")
    assert code == 0


def test_csv_datev(lauf, tmp_path):
    out = tmp_path / "b.csv"
    code, _ = lauf("csv", FIX / "buchungen.xlsx", "--format", "datev", "--out", out)
    assert code == 0
    roh = out.read_bytes()
    assert b"\r\n" in roh and not roh.startswith(b"\xef\xbb\xbf")
    text = roh.decode("cp1252")
    zeilen = text.split("\r\n")
    assert zeilen[0] == '"Belegdatum";"Konto";"Gegenkonto";"Betrag netto";"Betrag brutto";"Buchungstext"'
    assert zeilen[1] == '15.01.2026;8400;10001;1000;1190;"Rechnung 2026-001 Müller; Wartung"'
    assert zeilen[2] == '31.01.2026;8400;10002;2499,99;2974,99;"Projekt ""Fabric"" Phase 1"'
    assert zeilen[3] == '03.02.2026;4930;70001;89,9;;"Bürobedarf"'
    hinweise = json.loads((tmp_path / "b.aenderungen.json").read_text())["hinweise"]
    assert any(h["zelle"] == "E4" for h in hinweise)


def test_csv_excel_de_und_utf8(lauf, tmp_path):
    lauf("csv", FIX / "buchungen.xlsx", "--format", "excel-de", "--out", tmp_path / "d.csv")
    roh = (tmp_path / "d.csv").read_bytes()
    assert roh.startswith(b"\xef\xbb\xbf")
    assert roh.decode("utf-8-sig").splitlines()[3] == "03.02.2026;4930;70001;89,9;;Bürobedarf"
    lauf("csv", FIX / "buchungen.xlsx", "--format", "utf8", "--out", tmp_path / "u.csv")
    roh = (tmp_path / "u.csv").read_bytes()
    assert b"\r\n" not in roh
    assert roh.decode().splitlines()[2] == '2026-01-31,8400,10002,2499.99,2974.99,"Projekt ""Fabric"" Phase 1"'


def test_csv_datev_bricht_bei_nicht_ansi_ab(lauf, tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Datum", "Text", "Betrag"])
    ws.append(["01.01.2026", "Pfeil \u2192 rechts", 1])
    wb.save(tmp_path / "x.xlsx")
    code, erg = lauf("csv", tmp_path / "x.xlsx", "--format", "datev", "--out", tmp_path / "x.csv")
    assert code == 3 and not (tmp_path / "x.csv").exists()
    assert erg["details"][0]["zelle"] == "B2"


def test_xlsm_behaelt_vba(lauf, tmp_path):
    quelle = FIX / "makro.xlsm"
    vorher = sha(quelle)
    code, _ = lauf("bereinigen", quelle, "--out", tmp_path / "m.xlsx")
    assert code == 2
    code, _ = lauf("bereinigen", quelle, "--out", tmp_path / "m.xlsm")
    assert code == 0 and sha(quelle) == vorher
    with zipfile.ZipFile(quelle) as a, zipfile.ZipFile(tmp_path / "m.xlsm") as b:
        assert a.read("xl/vbaProject.bin") == b.read("xl/vbaProject.bin")
        assert "macroEnabled" in b.read("[Content_Types].xml").decode()
    ws = load_workbook(tmp_path / "m.xlsm").active
    assert ws["C2"].value == 2.5 and ws["B2"].value == "10"


def test_grenzen(lauf, tmp_path):
    (tmp_path / "alt.xls").write_bytes(b"x")
    code, erg = lauf("profil", tmp_path / "alt.xls", "--out", tmp_path / "r.md")
    assert code == 2 and ".xlsx" in erg["fehler"]
    (tmp_path / "kaputt.xlsx").write_bytes(b"kein zip")
    code, erg = lauf("profil", tmp_path / "kaputt.xlsx", "--out", tmp_path / "r.md")
    assert code == 2
    code, erg = lauf("profil", tmp_path / "fehlt.xlsx", "--out", tmp_path / "r.md")
    assert code == 2


def test_markdown_bleibt_heil(lauf, tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Ort", "Wert"])
    ws.append(["a | b", "Zeile\n# eingeschleust", 1])
    ws.append(["c", "d", "1.234"])
    wb.save(tmp_path / "x.xlsx")
    code, _ = lauf("profil", tmp_path / "x.xlsx", "--out", tmp_path / "r.md")
    assert code == 0
    md = (tmp_path / "r.md").read_text()
    assert not any(z.startswith("# eingeschleust") for z in md.splitlines())


def test_schema(lauf):
    code, erg = lauf("schema")
    assert code == 0 and erg["version"] == "0.1.1"
    assert erg["csv_formate"]["datev"]["kodierung"] == "cp1252"
