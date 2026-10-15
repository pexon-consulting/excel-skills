"""Dateien öffnen mit festen Grenzen. Nichts in diesem Modul schreibt in eine Eingabedatei."""

from __future__ import annotations

import warnings
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from . import regeln
from .regeln import ist_leer

MAX_BYTES = 50 * 1024 * 1024
MAX_ENTPACKT = 500 * 1024 * 1024
MAX_ZELLEN = 1_000_000
ENDUNGEN = (".xlsx", ".xlsm")


class Abbruch(Exception):
    """Fehler mit Exit-Code. 2: Aufruf oder Datei, 3: eine Regel verbietet das Schreiben."""

    def __init__(self, meldung: str, code: int = 2, details: list[Any] | None = None):
        super().__init__(meldung)
        self.meldung = meldung
        self.code = code
        self.details = details or []


def pruefe_eingabe(pfad: Path) -> Path:
    if not pfad.is_file():
        raise Abbruch(f"Datei nicht gefunden: {pfad}")
    if pfad.name.startswith("~$"):
        raise Abbruch(f"Sperrdatei von Excel, keine Arbeitsmappe: {pfad}")
    endung = pfad.suffix.lower()
    if endung == ".xls":
        raise Abbruch(f"{pfad.name}: altes .xls-Format. Bitte in Excel als .xlsx speichern.")
    if endung not in ENDUNGEN:
        raise Abbruch(f"{pfad.name}: nur .xlsx und .xlsm werden gelesen.")
    groesse = pfad.stat().st_size
    if groesse > MAX_BYTES:
        raise Abbruch(f"{pfad.name}: {groesse // (1024 * 1024)} MB, Grenze ist 50 MB.")
    try:
        with zipfile.ZipFile(pfad) as z:
            entpackt = sum(i.file_size for i in z.infolist())
    except zipfile.BadZipFile as exc:
        raise Abbruch(f"{pfad.name}: keine gültige Excel-Datei ({exc}).") from exc
    if entpackt > MAX_ENTPACKT:
        raise Abbruch(f"{pfad.name}: entpackt über 500 MB, Abbruch.")
    return pfad


def lade(pfad: Path, werte: bool = False, vba: bool = False) -> Workbook:
    """Öffnet die Mappe. werte=True liefert die von Excel gespeicherten Ergebnisse statt Formeln.

    vba=True hält das VBA-Projekt einer .xlsm für das Speichern einer Kopie fest. Danach `schliessen()` aufrufen.
    """
    pruefe_eingabe(pfad)
    mit_vba = vba and pfad.suffix.lower() == ".xlsm" and not werte
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        try:
            wb = load_workbook(pfad, data_only=werte, keep_vba=mit_vba)
        except Exception as exc:  # openpyxl wirft viele Typen bei kaputten Dateien
            raise Abbruch(f"{pfad.name}: lässt sich nicht öffnen ({type(exc).__name__}: {exc}).") from exc
    for ws in wb.worksheets:
        zellen = ws.max_row * ws.max_column
        if zellen > MAX_ZELLEN:
            raise Abbruch(f"{pfad.name}, Blatt {ws.title}: {zellen} Zellen, Grenze ist 1.000.000 je Blatt.")
    return wb


def schliessen(wb: Workbook) -> None:
    archiv = getattr(wb, "vba_archive", None)
    if archiv is not None:
        archiv.close()
        wb.vba_archive = None


def lade_beide(pfad: Path) -> tuple[Workbook, Workbook]:
    """Mappe mit Formeln und dieselbe Mappe mit den von Excel gespeicherten Ergebnissen."""
    return lade(pfad), lade(pfad, werte=True)


def ist_dieselbe(a: Path, b: Path) -> bool:
    """Auch bei anderer Groß- und Kleinschreibung oder per Hardlink dieselbe Datei."""
    if a.resolve() == b.resolve():
        return True
    try:
        return a.exists() and b.exists() and a.samefile(b)
    except OSError:
        return False


def pruefe_ausgabe(out: Path | None, eingaben: list[Path], endungen: tuple[str, ...]) -> Path:
    """--out ist Pflicht, hat die richtige Endung und ist keine Eingabedatei."""
    if out is None:
        raise Abbruch("--out fehlt. Das Original wird nie überschrieben, deshalb braucht jede Ausgabe einen eigenen Pfad.")
    if out.suffix.lower() not in endungen:
        raise Abbruch(f"--out muss auf {' oder '.join(endungen)} enden: {out}")
    for e in eingaben:
        if ist_dieselbe(out, e):
            raise Abbruch(f"--out ist die Eingabedatei {e}. Das Original wird nie überschrieben.")
    if out.exists() and out.is_dir():
        raise Abbruch(f"--out ist ein Ordner: {out}")
    out.parent.mkdir(parents=True, exist_ok=True)
    return out


def begleitdatei(out: Path, name: str) -> Path:
    """Pfad für `aenderungen.json` oder `report.json` neben einer Ausgabe."""
    return out.with_name(f"{out.stem}.{name}")


def blatt(wb: Workbook, name: str | None) -> Worksheet:
    """Ein Tabellenblatt. Diagrammblätter zählen nicht."""
    blaetter = {ws.title: ws for ws in wb.worksheets}
    if not blaetter:
        raise Abbruch("Die Mappe hat kein Tabellenblatt.")
    if name is None:
        return wb.worksheets[0]
    if name not in blaetter:
        raise Abbruch(f"Tabellenblatt {name!r} fehlt. Vorhanden: {', '.join(blaetter)}")
    return blaetter[name]


def ist_formel(zelle: Any) -> bool:
    return zelle.data_type == "f"


def formeltext(zelle: Any) -> str:
    """Formel als Text. Datentabellen aus der Was-wäre-wenn-Analyse haben keinen Formeltext."""
    v = zelle.value
    if isinstance(v, str):
        return v
    text = getattr(v, "text", None)
    return text if isinstance(text, str) else f"<{type(v).__name__}>"


@dataclass
class Raster:
    """Werte eines Blatts als Liste von Zeilen, je Zelle (wert, ist_formel)."""

    zeilen: list[list[tuple[Any, bool]]]
    spalten_genutzt: list[int]

    @property
    def breite(self) -> int:
        return len(self.spalten_genutzt)


def lese_raster(ws: Worksheet) -> Raster:
    zeilen: list[list[tuple[Any, bool]]] = []
    genutzt: set[int] = set()
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        zeile = []
        for c in row:
            f = ist_formel(c)
            zeile.append((c.value, f))
            if f or not ist_leer(c.value):
                genutzt.add(c.column)
        zeilen.append(zeile)
    return Raster(zeilen, sorted(genutzt))


@dataclass
class Tabelle:
    """Kopfzeile und Datenzeilen eines Blatts nach #kopfzeile."""

    blatt: str
    kopf_zeile: int | None  # Excel-Zeilennummer
    begruendung: str
    erste_spalte: int
    letzte_spalte: int
    kopf: list[str]
    zeilen: list[tuple[int, list[Any]]] = field(default_factory=list)  # (Excel-Zeile, Werte)

    def spalte(self, name: str) -> int:
        try:
            return self.kopf.index(name)
        except ValueError as exc:
            raise Abbruch(f"Spalte {name!r} fehlt in Blatt {self.blatt}. Vorhanden: {', '.join(self.kopf)}") from exc


def kopfname(wert: Any, spalte: int) -> str:
    if ist_leer(wert):
        return f"Spalte {get_column_letter(spalte)}"
    return str(wert).strip(regeln.LEER)


def tabelle_pflicht(ws: Worksheet, werte_ws: Worksheet | None = None, datei: Path | None = None) -> Tabelle:
    """Wie tabelle(), bricht aber mit Exit-Code 3 ab, wenn keine Kopfzeile gefunden wird."""
    t = tabelle(ws, werte_ws)
    if t.kopf_zeile is None:
        wo = f"{datei.name}, Blatt {ws.title}" if datei else f"Blatt {ws.title}"
        raise Abbruch(f"{wo}: keine Kopfzeile gefunden ({t.begruendung}). Nichts geschrieben.", code=3)
    return t


def tabelle(ws: Worksheet, werte_ws: Worksheet | None = None, raster: Raster | None = None) -> Tabelle:
    """Liest die Tabelle ab der Kopfzeile. Formelzellen nehmen den gespeicherten Wert aus werte_ws, sonst den Formeltext."""
    r = raster or lese_raster(ws)
    if not r.spalten_genutzt:
        return Tabelle(ws.title, None, "Blatt ist leer", 1, 0, [])
    links, rechts = r.spalten_genutzt[0], r.spalten_genutzt[-1]
    idx, grund = regeln.finde_kopfzeile([z[links - 1 : rechts] for z in r.zeilen], r.breite)
    if idx is None:
        return Tabelle(ws.title, None, grund, links, rechts, [])
    kopf = [kopfname(w, links + i) for i, (w, _) in enumerate(r.zeilen[idx][links - 1 : rechts])]
    t = Tabelle(ws.title, idx + 1, grund, links, rechts, kopf)
    for nr, zeile in enumerate(r.zeilen[idx + 1 :], start=idx + 2):
        teil = zeile[links - 1 : rechts]
        if all(ist_leer(w) and not f for w, f in teil):
            continue
        werte = []
        for off, (w, f) in enumerate(teil):
            if f and werte_ws is not None:
                cache = werte_ws.cell(row=nr, column=links + off).value
                werte.append(cache if cache is not None else (w if isinstance(w, str) else formeltext(ws.cell(row=nr, column=links + off))))
            else:
                werte.append(w)
        t.zeilen.append((nr, werte))
    return t
