"""excel-skills profil: Kopfzeile, Datenbereich, verbundene Zellen, Summenzeilen, Spaltentypen, Urteil."""

from __future__ import annotations

from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import bericht, laden, regeln


def hinweise_zelle(wert: Any, ref: str) -> dict[str, Any] | None:
    if not isinstance(wert, str) or wert.startswith("="):
        return None
    if regeln.deutsche_zahl(wert).mehrdeutig:
        return {"zelle": ref, "regel": "deutsche-zahl", "wert": wert, "hinweis": "ohne Komma mehrdeutig, bleibt Text"}
    if regeln.ist_us_datum(wert):
        return {"zelle": ref, "regel": "datum", "wert": wert, "hinweis": "Datum mit Schrägstrich, Reihenfolge Tag und Monat unklar, bleibt Text"}
    return None


def blatt_profil(ws: Worksheet) -> dict[str, Any]:
    r = laden.raster(ws)
    t = laden.tabelle(ws)
    verbunden = sorted((str(m) for m in ws.merged_cells.ranges), key=lambda s: (len(s), s))
    erg: dict[str, Any] = {
        "blatt": ws.title,
        "kopfzeile": {"zeile": t.kopf_zeile, "begruendung": t.begruendung},
        "datenbereich": None,
        "verbundene_zellen": verbunden,
        "verbundene_im_datenbereich": [],
        "summenzeilen": [],
        "spalten": [],
        "urteil": "formular",
        "gruende": [],
        "hinweise": [],
    }
    if r.breite < regeln.FORMULAR_MAX_SPALTEN:
        erg["gruende"].append(f"nur {r.breite} Spalten mit Daten, weniger als {regeln.FORMULAR_MAX_SPALTEN}")
        return erg
    if t.kopf_zeile is None:
        erg["urteil"] = "bericht"
        erg["gruende"].append(t.begruendung)
        return erg

    datenzeilen = [nr for nr, _ in t.zeilen]
    if datenzeilen:
        erg["datenbereich"] = f"{get_column_letter(t.erste_spalte)}{datenzeilen[0]}:{get_column_letter(t.letzte_spalte)}{datenzeilen[-1]}"
    erg["verbundene_im_datenbereich"] = [str(m) for m in ws.merged_cells.ranges if m.min_row > t.kopf_zeile]
    summen = [nr for nr, werte in t.zeilen if regeln.ist_summenzeile(werte)]
    erg["summenzeilen"] = summen

    for off, name in enumerate(t.kopf):
        spalte = t.erste_spalte + off
        typen: Counter[str] = Counter()
        for nr in datenzeilen:
            if nr in summen:
                continue
            w, f = r.zeilen[nr - 1][spalte - 1]
            typen[regeln.zelltyp(w, f)] += 1
            h = hinweise_zelle(w, f"{get_column_letter(spalte)}{nr}")
            if h:
                erg["hinweise"].append(h)
        gefuellt = sum(n for k, n in typen.items() if k != "leer")
        gesamt = sum(typen.values())
        typ = "leer"
        if gefuellt:
            top, n = max(((k, n) for k, n in typen.items() if k != "leer"), key=lambda x: (x[1], x[0]))
            typ = top if n / gefuellt >= regeln.SPALTENTYP_ANTEIL else "gemischt"
        erg["spalten"].append({
            "spalte": get_column_letter(spalte),
            "name": name,
            "typ": typ,
            "leer_anteil": round(typen["leer"] / gesamt, 2) if gesamt else 0.0,
            "typen": dict(sorted(typen.items())),
        })

    if erg["verbundene_im_datenbereich"]:
        erg["gruende"].append(f"{len(erg['verbundene_im_datenbereich'])} verbundene Bereiche im Datenbereich")
    mitten = [nr for nr in summen if datenzeilen and nr < datenzeilen[-1]]
    if mitten:
        erg["gruende"].append(f"Summenzeile vor der letzten Datenzeile: {', '.join(map(str, mitten))}")
    if t.kopf_zeile > 1:
        erg["gruende"].append(f"Kopfzeile erst in Zeile {t.kopf_zeile}")
    erg["urteil"] = "bericht" if erg["verbundene_im_datenbereich"] or mitten else "tabelle"
    if erg["urteil"] == "tabelle" and t.kopf_zeile == 1:
        erg["gruende"].append("Kopfzeile in Zeile 1, keine verbundenen Zellen, keine Summenzeile mitten drin")
    elif erg["urteil"] == "tabelle":
        erg["gruende"].append("keine verbundenen Zellen und keine Summenzeile mitten drin")
    return erg


def ausfuehren(datei: Path, out: Path) -> dict[str, Any]:
    wb = laden.lade(datei)
    blaetter = [blatt_profil(ws) for ws in wb.worksheets]
    daten = {"befehl": "profil", "datei": datei.name, "blaetter": blaetter}
    zeilen: list[str] = [f"Datei: `{bericht.zelle_md(datei.name)}`", ""]
    for b in blaetter:
        kz = b["kopfzeile"]
        zeilen += [
            f"## Blatt {bericht.zelle_md(b['blatt'])}",
            "",
            f"- Urteil: **{b['urteil']}**",
            f"- Kopfzeile: {kz['zeile'] if kz['zeile'] else 'keine'} ({kz['begruendung']})",
            f"- Datenbereich: {b['datenbereich'] or 'keiner'}",
            f"- Verbundene Zellen: {len(b['verbundene_zellen'])}, davon im Datenbereich {len(b['verbundene_im_datenbereich'])}",
            f"- Summenzeilen: {', '.join(map(str, b['summenzeilen'])) or 'keine'}",
            f"- Gründe: {'; '.join(b['gruende'])}",
            "",
        ]
        zeilen += bericht.tabelle_md(
            ["Spalte", "Name", "Typ", "leer"],
            [[s["spalte"], s["name"], s["typ"], f"{round(s['leer_anteil'] * 100)} %"] for s in b["spalten"]],
        )
        if b["hinweise"]:
            zeilen += ["Hinweise:", ""]
            zeilen += bericht.tabelle_md(["Zelle", "Wert", "Hinweis"], [[h["zelle"], h["wert"], h["hinweis"]] for h in b["hinweise"]])
    json_pfad = bericht.schreibe_bericht(out, f"Excel-Profil: {bericht.zelle_md(datei.name)}", zeilen, daten)
    return {
        "befehl": "profil",
        "bericht": str(out),
        "json": str(json_pfad),
        "blaetter": {b["blatt"]: b["urteil"] for b in blaetter},
    }
