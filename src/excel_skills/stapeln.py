"""excel-skills stapeln: mehrere Excel-Dateien mit gleichen Spalten zu einer Tabelle zusammenführen."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formula.translate import Translator

from . import bericht, laden, regeln

QUELLSPALTE = "quelle_datei"


def dateien(ordner: Path, out: Path) -> list[Path]:
    if not ordner.is_dir():
        raise laden.Abbruch(f"Ordner nicht gefunden: {ordner}")
    liste = sorted(
        p for p in ordner.iterdir()
        if p.is_file() and p.suffix.lower() in laden.ENDUNGEN and not p.name.startswith("~$") and not laden.ist_dieselbe(p, out)
    )
    if not liste:
        raise laden.Abbruch(f"Keine .xlsx oder .xlsm in {ordner}")
    return liste


def ausfuehren(ordner: Path, blatt: str | None, out: Path | None, werte: bool) -> dict[str, Any]:
    if out is None:
        raise laden.Abbruch("--out fehlt. Das Original wird nie überschrieben, deshalb braucht jede Ausgabe einen eigenen Pfad.")
    liste = dateien(ordner, out)
    out = laden.pruefe_ausgabe(out, liste, (".xlsx",))

    gelesen = []
    for p in liste:
        wb = laden.lade(p)
        ws = laden.blatt(wb, blatt)
        wv = laden.lade(p, werte=True)[ws.title] if werte else None
        t = laden.tabelle_pflicht(ws, datei=p)
        gelesen.append((p, ws, wv, t))

    basis = gelesen[0][3].kopf
    abweichungen = []
    for p, _, _, t in gelesen:
        doppelt = sorted({k for k in t.kopf if t.kopf.count(k) > 1})
        fehlt = [k for k in basis if k not in t.kopf]
        extra = [k for k in t.kopf if k not in basis]
        if fehlt or extra or doppelt:
            abweichungen.append({"datei": p.name, "fehlt": fehlt, "zusaetzlich": extra, "doppelt": doppelt})
    if abweichungen:
        raise laden.Abbruch(
            f"Spalten passen nicht zu {gelesen[0][0].name}. Nichts geschrieben.", code=3, details=abweichungen
        )
    if QUELLSPALTE in basis:
        raise laden.Abbruch(f"Spalte {QUELLSPALTE} gibt es schon in den Eingaben.", code=3)

    ziel = Workbook()
    zs = ziel.active
    zs.title = "Gesamt"
    zs.append([*basis, QUELLSPALTE])
    uebersprungen: list[dict[str, Any]] = []
    konflikte: list[dict[str, Any]] = []
    ohne_cache: list[dict[str, Any]] = []
    quellen = []
    zeile_ziel = 2
    for p, ws, wv, t in gelesen:
        anzahl = 0
        for nr, werte_zeile in t.zeilen:
            if regeln.ist_summenzeile(werte_zeile):
                uebersprungen.append({"datei": p.name, "zeile": nr, "regel": "summen",
                                      "werte": [bericht.jsonwert(w) for w in werte_zeile if not laden.ist_leer(w)]})
                continue
            for ziel_sp, name in enumerate(basis, start=1):
                quell_sp = t.erste_spalte + t.kopf.index(name)
                c = ws.cell(row=nr, column=quell_sp)
                zelle = zs.cell(row=zeile_ziel, column=ziel_sp)
                if laden.ist_formel(c):
                    text = laden.formeltext(c)
                    if wv is not None:
                        cache = wv.cell(row=nr, column=quell_sp).value
                        if cache is None:
                            ohne_cache.append({"datei": p.name, "zelle": c.coordinate, "formel": text})
                        zelle.value = cache
                    elif "!" in text or "[" in text:
                        konflikte.append({"datei": p.name, "zelle": c.coordinate, "formel": text,
                                          "grund": "Formel verweist auf ein anderes Blatt oder eine andere Datei"})
                    elif t.kopf != basis:
                        konflikte.append({"datei": p.name, "zelle": c.coordinate, "formel": text,
                                          "grund": "Spalten stehen in anderer Reihenfolge, die Bezüge der Formel würden falsch"})
                    else:
                        zelle.value = Translator(text, origin=c.coordinate).translate_formula(zelle.coordinate)
                else:
                    zelle.value = c.value
                zelle.number_format = c.number_format
            zs.cell(row=zeile_ziel, column=len(basis) + 1).value = p.name
            zeile_ziel += 1
            anzahl += 1
        quellen.append({"datei": p.name, "blatt": ws.title, "kopfzeile": t.kopf_zeile, "zeilen": anzahl})
    if konflikte:
        raise laden.Abbruch(
            "Formeln lassen sich nicht sicher mitnehmen. Nichts geschrieben. "
            "Mit --werte werden die gespeicherten Ergebnisse übernommen.", code=3, details=konflikte,
        )

    ziel.save(out)
    protokoll = laden.begleitdatei(out, "aenderungen.json")
    hinweise = [{"zelle": f"{h['datei']}!{h['zelle']}", "hinweis": "Formel ohne gespeichertes Ergebnis, Zelle bleibt leer"} for h in ohne_cache]
    if werte:
        hinweise.insert(0, {"zelle": None, "hinweis": "--werte: Formeln wurden als gespeicherte Ergebnisse übernommen"})
    daten = {
        "befehl": "stapeln",
        "ordner": ordner.name,
        "ausgabe": out.name,
        "spalten": [*basis, QUELLSPALTE],
        "quellen": quellen,
        "zeilen_gesamt": zeile_ziel - 2,
        "uebersprungen": uebersprungen,
        "hinweise": hinweise,
    }
    bericht.schreibe_json(protokoll, daten)
    return {
        "befehl": "stapeln",
        "ausgabe": str(out),
        "aenderungen": str(protokoll),
        "dateien": len(quellen),
        "zeilen_gesamt": zeile_ziel - 2,
        "uebersprungen": len(uebersprungen),
    }
