"""excel-skills bereinigen: schreibt eine bereinigte Kopie. Das Original bleibt unverändert."""

from __future__ import annotations

import bisect
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.worksheet import Worksheet

from . import bericht, laden, regeln

REGELN = ("kopfzeile", "verbunden", "leerzeilen", "summen", "zahlen", "datum", "trimmen")
_ZELLE = re.compile(r"^(\$?[A-Za-z]{1,3})(\$?)(\d+)$")
_ZEILEN = re.compile(r"^(\$?)(\d+):(\$?)(\d+)$")


class Zeilenplan:
    """Welche Zeilen eines Blatts entfallen, und wohin die übrigen rutschen."""

    def __init__(self, blatt: str, geloescht: set[int]):
        self.blatt = blatt
        self.geloescht = sorted(geloescht)
        self._menge = set(geloescht)

    def neu(self, zeile: int) -> int | None:
        if zeile in self._menge:
            return None
        return zeile - bisect.bisect_left(self.geloescht, zeile)

    def bereich(self, von: int, bis: int) -> tuple[int, int] | None:
        a = von - bisect.bisect_left(self.geloescht, von)
        b = bis - bisect.bisect_right(self.geloescht, bis)
        return (a, b) if a <= b else None


def _blattname(prefix: str) -> str | None:
    if "[" in prefix:
        return None
    if prefix.startswith("'") and prefix.endswith("'"):
        return prefix[1:-1].replace("''", "'")
    return prefix


def verschiebe_bezug(ref: str, plan: Zeilenplan) -> str | None:
    """A5, A5:B9 oder 5:9 auf die neuen Zeilen. None, wenn der Bezug ganz entfällt."""
    m = _ZELLE.match(ref)
    if m:
        neu = plan.neu(int(m.group(3)))
        return None if neu is None else f"{m.group(1)}{m.group(2)}{neu}"
    if ":" in ref:
        a, b = ref.split(":", 1)
        ma, mb = _ZELLE.match(a), _ZELLE.match(b)
        if ma and mb:
            r = plan.bereich(int(ma.group(3)), int(mb.group(3)))
            return None if r is None else f"{ma.group(1)}{ma.group(2)}{r[0]}:{mb.group(1)}{mb.group(2)}{r[1]}"
        mz = _ZEILEN.match(ref)
        if mz:
            r = plan.bereich(int(mz.group(2)), int(mz.group(4)))
            return None if r is None else f"{mz.group(1)}{r[0]}:{mz.group(3)}{r[1]}"
    return ref


def verschiebe_formel(formel: str, auf_blatt: str, plan: Zeilenplan) -> tuple[str, str | None]:
    """Passt Zeilenbezüge auf plan.blatt an. Rückgabe: neue Formel und Konfliktgrund oder None."""
    try:
        tok = Tokenizer(formel)
    except Exception as exc:
        return formel, f"Formel nicht lesbar: {exc}"
    geaendert = False
    for item in tok.items:
        if item.type != "OPERAND" or item.subtype != "RANGE":
            continue
        if "!" in item.value:
            prefix, ref = item.value.rsplit("!", 1)
            if _blattname(prefix) != plan.blatt:
                continue
        else:
            prefix, ref = None, item.value
            if auf_blatt != plan.blatt:
                continue
        neu = verschiebe_bezug(ref, plan)
        if neu is None:
            return formel, f"Bezug {item.value} zeigt auf eine Zeile, die entfällt"
        if neu != ref:
            item.value = f"{prefix}!{neu}" if prefix is not None else neu
            geaendert = True
    return (tok.render() if geaendert else formel), None


@dataclass
class Lauf:
    """Alles, was ein Lauf sammelt, bevor er schreibt."""

    ws: Worksheet
    aktiv: list[str]
    kopf: int
    geloescht: dict[int, str] = field(default_factory=dict)
    formeln: dict[tuple[str, str], Any] = field(default_factory=dict)
    namen: list[tuple[Any, str]] = field(default_factory=list)
    bereiche: list[tuple[Any, str, str]] = field(default_factory=list)  # (Objekt, Attribut, neuer Bezug)
    druckbereich: str | None = None
    verbinden: list[tuple[int, int, int, int]] = field(default_factory=list)
    konflikte: list[dict[str, Any]] = field(default_factory=list)
    aenderungen: list[dict[str, Any]] = field(default_factory=list)
    hinweise: list[dict[str, Any]] = field(default_factory=list)

    def aenderung(self, zelle: str, alt: Any, neu: Any, regel: str) -> None:
        self.aenderungen.append({"blatt": self.ws.title, "zelle": zelle, "alt": alt, "neu": neu, "regel": regel})

    def konflikt(self, blatt: str | None, zelle: str, formel: str, grund: str) -> None:
        self.konflikte.append({"blatt": blatt, "zelle": zelle, "formel": formel, "grund": grund})


def _zeilen_planen(lauf: Lauf, t: laden.Tabelle) -> None:
    ws, kopf = lauf.ws, lauf.kopf
    if "kopfzeile" in lauf.aktiv:
        lauf.geloescht.update({nr: "kopfzeile" for nr in range(1, kopf)})
    letzte = t.zeilen[-1][0] if t.zeilen else kopf
    if "leerzeilen" in lauf.aktiv:
        for nr in range(kopf + 1, letzte):
            if all(laden.ist_leer(c.value) and not laden.ist_formel(c) for c in ws[nr]):
                lauf.geloescht[nr] = "leerzeilen"
    if "summen" in lauf.aktiv:
        lauf.geloescht.update({nr: "summen" for nr, werte in t.zeilen if regeln.ist_summenzeile(werte)})


def _bezuege_planen(lauf: Lauf, wb: Workbook, plan: Zeilenplan) -> None:
    """Neue Formeln, Namen und Bereiche berechnen. Schreibt noch nichts."""
    ws = lauf.ws
    for fws in wb.worksheets:
        for row in fws.iter_rows():
            for c in row:
                if not laden.ist_formel(c) or (fws is ws and c.row in lauf.geloescht):
                    continue
                v = c.value
                if not isinstance(v, (str, ArrayFormula)):
                    lauf.konflikt(fws.title, c.coordinate, laden.formeltext(c), "Datentabelle (Was-wäre-wenn) wird nicht verschoben")
                    continue
                text = laden.formeltext(c)
                neu, grund = verschiebe_formel(text, fws.title, plan)
                if grund:
                    lauf.konflikt(fws.title, c.coordinate, text, grund)
                    continue
                if isinstance(v, ArrayFormula):
                    ref = (verschiebe_bezug(v.ref, plan) or v.ref) if fws is ws else v.ref
                    if neu != text or ref != v.ref:
                        lauf.formeln[(fws.title, c.coordinate)] = ArrayFormula(ref=ref, text=neu)
                elif neu != text:
                    lauf.formeln[(fws.title, c.coordinate)] = neu

    namensraeume = [("Arbeitsmappe", wb.defined_names, "")] + [(s.title, s.defined_names, s.title) for s in wb.worksheets]
    for wo, namen, blatt in namensraeume:
        for name, dn in namen.items():
            text = dn.attr_text or ""
            neu, grund = verschiebe_formel("=" + text, blatt, plan)
            if grund:
                lauf.konflikt(wo, f"Name {name}", text, grund)
            elif neu != "=" + text:
                lauf.namen.append((dn, neu[1:]))

    for tab in ws.tables.values():
        r1 = int(_ZELLE.match(tab.ref.split(":")[0]).group(3))
        neu = verschiebe_bezug(tab.ref, plan)
        if r1 in lauf.geloescht or neu is None:
            lauf.konflikt(ws.title, f"Tabelle {tab.displayName}", tab.ref, "Kopfzeile der intelligenten Tabelle entfällt")
            continue
        lauf.bereiche.append((tab, "ref", neu))
        if tab.autoFilter is not None and tab.autoFilter.ref:
            lauf.bereiche.append((tab.autoFilter, "ref", neu))
    if ws.auto_filter.ref:
        neu = verschiebe_bezug(ws.auto_filter.ref, plan)
        if neu is None:
            lauf.hinweise.append({"zelle": ws.auto_filter.ref, "regel": "zeilen", "hinweis": "Autofilter entfällt mit seinen Zeilen"})
            neu = ""
        lauf.bereiche.append((ws.auto_filter, "ref", neu))
    if ws.print_area:
        teile = [verschiebe_bezug(p.rsplit("!", 1)[-1].replace("$", ""), plan) for p in ws.print_area.split(",")]
        if all(teile):
            lauf.druckbereich = ",".join(teile)  # type: ignore[arg-type]
        else:
            lauf.hinweise.append({"zelle": ws.print_area, "regel": "zeilen", "hinweis": "Druckbereich entfällt mit seinen Zeilen"})
            lauf.druckbereich = ""


def _verbundene_aufloesen(lauf: Lauf, plan: Zeilenplan) -> None:
    ws = lauf.ws
    for m in sorted(ws.merged_cells.ranges, key=lambda m: (m.min_row, m.min_col)):
        bereich = str(m)
        r1, c1, r2, c2 = m.min_row, m.min_col, m.max_row, m.max_col
        oben_links = ws.cell(row=r1, column=c1)
        wert = oben_links.value
        ws.unmerge_cells(bereich)
        if all(r in lauf.geloescht for r in range(r1, r2 + 1)):
            continue
        if "verbunden" not in lauf.aktiv or r2 < lauf.kopf:
            rb = plan.bereich(r1, r2)
            if rb:
                lauf.verbinden.append((rb[0], c1, rb[1], c2))
            continue
        if laden.ist_formel(oben_links):
            lauf.hinweise.append({"zelle": bereich, "regel": "verbunden", "hinweis": "verbunden mit Formel, nur die Zelle oben links behält die Formel"})
            lauf.aenderung(bereich, "verbunden", "aufgelöst", "verbunden")
            continue
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                if (r, c) != (r1, c1):
                    ws.cell(row=r, column=c).value = wert
        lauf.aenderung(bereich, "verbunden", wert, "verbunden")


def _werte_umwandeln(lauf: Lauf, t: laden.Tabelle) -> None:
    ws = lauf.ws
    for nr in range(lauf.kopf, ws.max_row + 1):
        if nr in lauf.geloescht:
            continue
        for sp in range(t.erste_spalte, t.letzte_spalte + 1):
            c = ws.cell(row=nr, column=sp)
            v = c.value
            if not isinstance(v, str) or laden.ist_formel(c):
                continue
            if "trimmen" in lauf.aktiv:
                s = v.strip(regeln.LEER)
                if s != v:
                    c.value = s
                    lauf.aenderung(c.coordinate, v, s, "trimmen")
                    v = s
            if nr == lauf.kopf:
                continue
            z = regeln.deutsche_zahl(v) if "zahlen" in lauf.aktiv else None
            if z is not None and z.wert is not None:
                c.value = z.wert
                if c.number_format == "@":
                    c.number_format = "General"
                lauf.aenderung(c.coordinate, v, z.wert, "deutsche-zahl")
                continue
            d = regeln.deutsches_datum(v) if "datum" in lauf.aktiv else None
            if d is not None:
                c.value = d
                c.number_format = "DD.MM.YYYY"
                lauf.aenderung(c.coordinate, v, d, "datum")
                continue
            h = regeln.hinweis(v, c.coordinate)
            if h:
                lauf.hinweise.append(h)


def _anwenden(lauf: Lauf, wb: Workbook, plan: Zeilenplan) -> None:
    ws = lauf.ws
    for nr in sorted(lauf.geloescht, reverse=True):
        ws.delete_rows(nr, 1)
    for (titel, koord), formel in lauf.formeln.items():
        zelle = wb[titel][koord]
        if wb[titel] is ws:
            zelle = ws.cell(row=plan.neu(zelle.row), column=zelle.column)  # type: ignore[arg-type]
        zelle.value = formel
    for dn, text in lauf.namen:
        dn.attr_text = text
    for obj, attr, neu in lauf.bereiche:
        setattr(obj, attr, neu or None)
    if lauf.druckbereich is not None:
        ws.print_area = lauf.druckbereich or None
    for r1, c1, r2, c2 in lauf.verbinden:
        if (r1, c1) != (r2, c2):
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    for row in ws.iter_rows():
        for c in row:
            if c.hyperlink is not None:
                c.hyperlink.ref = c.coordinate


def ausfuehren(datei: Path, blatt_name: str | None, out: Path | None, ohne: set[str]) -> dict[str, Any]:
    unbekannt = ohne - set(REGELN)
    if unbekannt:
        raise laden.Abbruch(f"--ohne kennt nur {', '.join(REGELN)}, nicht {', '.join(sorted(unbekannt))}")
    xlsm = datei.suffix.lower() == ".xlsm"
    out = laden.pruefe_ausgabe(out, [datei], (".xlsm",) if xlsm else (".xlsx",))
    wb = laden.lade(datei, vba=True)
    try:
        return _bereinigen(wb, datei, blatt_name, out, ohne)
    finally:
        laden.schliessen(wb)


def _bereinigen(wb: Workbook, datei: Path, blatt_name: str | None, out: Path, ohne: set[str]) -> dict[str, Any]:
    ws = laden.blatt(wb, blatt_name)
    t = laden.tabelle_pflicht(ws)
    lauf = Lauf(ws, [r for r in REGELN if r not in ohne], t.kopf_zeile)  # type: ignore[arg-type]
    _zeilen_planen(lauf, t)
    plan = Zeilenplan(ws.title, set(lauf.geloescht))
    if lauf.geloescht:
        _bezuege_planen(lauf, wb, plan)
    if lauf.konflikte:
        raise laden.Abbruch(
            "Formeln oder Bereiche zeigen auf Zeilen, die entfallen würden. Nichts geschrieben. "
            "Mit --ohne summen oder --ohne leerzeilen erneut versuchen.",
            code=3, details=lauf.konflikte,
        )
    _verbundene_aufloesen(lauf, plan)
    _werte_umwandeln(lauf, t)
    if lauf.geloescht:
        for name, vorhanden in (("Datenüberprüfung", ws.data_validations.dataValidation), ("bedingte Formatierung", len(ws.conditional_formatting))):
            if vorhanden:
                lauf.hinweise.append({"zelle": None, "regel": "zeilen", "hinweis": f"{name} vorhanden, deren Bereich wurde nicht verschoben, bitte prüfen"})
    geloeschte_zeilen = [
        {"zeile": nr, "regel": lauf.geloescht[nr], "werte": [bericht.jsonwert(c.value) for c in ws[nr] if not laden.ist_leer(c.value)]}
        for nr in sorted(lauf.geloescht)
    ]
    _anwenden(lauf, wb, plan)

    wb.save(out)
    protokoll = laden.begleitdatei(out, "aenderungen.json")
    bericht.schreibe_json(protokoll, {
        "befehl": "bereinigen",
        "datei": datei.name,
        "blatt": ws.title,
        "ausgabe": out.name,
        "zellbezuege": "Zellen und Zeilen in der Zählung des Originals",
        "regeln_aktiv": lauf.aktiv,
        "kopfzeile_original": lauf.kopf,
        "geloeschte_zeilen": geloeschte_zeilen,
        "formeln_angepasst": len(lauf.formeln),
        "aenderungen": lauf.aenderungen,
        "hinweise": lauf.hinweise,
    })
    return {
        "befehl": "bereinigen",
        "ausgabe": str(out),
        "aenderungen": str(protokoll),
        "zellen_geaendert": len(lauf.aenderungen),
        "zeilen_geloescht": len(geloeschte_zeilen),
        "formeln_angepasst": len(lauf.formeln),
        "hinweise": len(lauf.hinweise),
        "kopfzeile_jetzt": 1 if "kopfzeile" in lauf.aktiv else lauf.kopf,
        "spalte_letzte": get_column_letter(t.letzte_spalte),
    }
