"""excel-skills csv: ein Blatt als CSV für DATEV, deutsches Excel oder UTF-8 ausgeben."""

from __future__ import annotations

import datetime as dt
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl.utils import get_column_letter

from . import bericht, laden, regeln

_ZEILENENDE = {"CRLF": "\r\n", "LF": "\n"}


def _zahl(x: float | int, dez: str) -> str:
    if isinstance(x, int):
        return str(x)
    if x.is_integer() and abs(x) < 1e15:
        return str(int(x))
    s = format(Decimal(repr(x)), "f")
    return s.replace(".", dez)


def _datum(d: dt.date, fmt: str) -> str:
    return d.strftime("%d.%m.%Y") if fmt == "TT.MM.JJJJ" else d.strftime("%Y-%m-%d")


def formatiere(wert: Any, f: dict[str, Any]) -> tuple[str, bool]:
    """Ein Zellwert als CSV-Text. Zweiter Wert: ist das ein Textfeld."""
    if wert is None:
        return "", False
    if isinstance(wert, bool):
        return f["wahrheitswerte"][0 if wert else 1], False
    if isinstance(wert, (int, float)):
        return _zahl(wert, f["dezimalzeichen"]), False
    if isinstance(wert, dt.datetime):
        s = _datum(wert.date(), f["datum"])
        if wert.time() != dt.time():
            s += " " + wert.strftime("%H:%M:%S")
        return s, False
    if isinstance(wert, dt.date):
        return _datum(wert, f["datum"]), False
    if isinstance(wert, dt.time):
        return wert.strftime("%H:%M:%S"), False
    return str(wert), True


def _feld(text: str, ist_text: bool, f: dict[str, Any]) -> str:
    noetig = any(z in text for z in (f["trennzeichen"], '"', "\r", "\n"))
    if noetig or (ist_text and f["text_in_anfuehrungszeichen"]):
        return '"' + text.replace('"', '""') + '"'
    return text


def ausfuehren(datei: Path, blatt: str | None, out: Path | None, fmt: str) -> dict[str, Any]:
    if fmt not in regeln.CSV_FORMATE:
        raise laden.Abbruch(f"--format kennt nur {', '.join(regeln.CSV_FORMATE)}")
    out = laden.pruefe_ausgabe(out, [datei], (".csv",))
    f = regeln.CSV_FORMATE[fmt]
    wb, wv = laden.lade_beide(datei)
    ws = laden.blatt(wb, blatt)
    vs = wv[ws.title]
    t = laden.tabelle_pflicht(ws)

    hinweise: list[dict[str, Any]] = []
    nicht_kodierbar: list[dict[str, Any]] = []
    formeln = 0
    zeilen_text: list[str] = []

    def zeile(werte: list[tuple[Any, str]]) -> None:
        felder = []
        for wert, ref in werte:
            text, ist_text = formatiere(wert, f)
            if ist_text:
                if f["kodierung"] == "cp1252":
                    try:
                        text.encode("cp1252")
                    except UnicodeEncodeError:
                        nicht_kodierbar.append({"zelle": ref, "wert": text})
                if text[:1] in ("=", "+", "-", "@") and ref:
                    hinweise.append({"zelle": ref, "hinweis": "Text beginnt mit = + - oder @, Tabellenprogramme können das als Formel lesen"})
            felder.append(_feld(text, ist_text, f))
        zeilen_text.append(f["trennzeichen"].join(felder))

    zeile([(k, f"{get_column_letter(t.erste_spalte + i)}{t.kopf_zeile}") for i, k in enumerate(t.kopf)])
    for nr, werte in t.zeilen:
        if regeln.ist_summenzeile(werte):
            hinweise.append({"zelle": f"Zeile {nr}", "hinweis": "Summenzeile wird mit ausgegeben, vorher bereinigen, wenn sie nicht in den Import soll"})
        reihe = []
        for off in range(len(t.kopf)):
            sp = t.erste_spalte + off
            c = ws.cell(row=nr, column=sp)
            wert = c.value
            if laden.ist_formel(c):
                formeln += 1
                wert = vs.cell(row=nr, column=sp).value
                if wert is None:
                    hinweise.append({"zelle": c.coordinate, "hinweis": "Formel ohne gespeichertes Ergebnis, Feld bleibt leer. Datei einmal in Excel öffnen und speichern."})
            if isinstance(wert, str) and wert in regeln.FEHLERWERTE:
                hinweise.append({"zelle": c.coordinate, "hinweis": f"Fehlerwert {wert} ({regeln.FEHLERWERTE[wert]}) wird als Text ausgegeben"})
            reihe.append((wert, c.coordinate))
        zeile(reihe)

    if nicht_kodierbar:
        raise laden.Abbruch(
            "Zeichen passen nicht in Windows-1252, das DATEV-Format braucht. Nichts geschrieben. "
            "Mit --format excel-de oder utf8 bleibt alles erhalten.", code=3, details=nicht_kodierbar,
        )
    ende = _ZEILENENDE[f["zeilenende"]]
    out.write_bytes((ende.join(zeilen_text) + ende).encode(f["kodierung"]))
    if formeln:
        hinweise.insert(0, {"zelle": None, "hinweis": f"{formeln} Formeln als gespeicherte Werte ausgegeben, CSV kennt keine Formeln"})
    protokoll = laden.begleitdatei(out, "aenderungen.json")
    daten = {
        "befehl": "csv",
        "datei": datei.name,
        "blatt": ws.title,
        "ausgabe": out.name,
        "format": fmt,
        "regeln": f,
        "kopfzeile_original": t.kopf_zeile,
        "spalten": t.kopf,
        "datenzeilen": len(t.zeilen),
        "hinweise": hinweise,
    }
    bericht.schreibe_json(protokoll, daten)
    return {
        "befehl": "csv",
        "ausgabe": str(out),
        "aenderungen": str(protokoll),
        "format": fmt,
        "datenzeilen": len(t.zeilen),
        "hinweise": len(hinweise),
    }
