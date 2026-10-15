"""excel-skills formeln: Fehlerwerte, harte Werte in Formelspalten, externe Verknüpfungen, Formeln, die nicht mitlaufen."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from . import bericht, laden, regeln

_EXTERN = re.compile(r"\[\d+\]")
_SUCHE = re.compile(r"\b(VLOOKUP|HLOOKUP|XLOOKUP|_xlfn\.XLOOKUP|MATCH|XMATCH|_xlfn\.XMATCH|INDEX)\s*\(", re.IGNORECASE)
_DEUTSCH = {"VLOOKUP": "SVERWEIS", "HLOOKUP": "WVERWEIS", "XLOOKUP": "XVERWEIS", "MATCH": "VERGLEICH", "XMATCH": "XVERGLEICH", "INDEX": "INDEX"}


def _hinweis(fehler: str, formel: str | None) -> str:
    if fehler == "#N/A" and formel:
        m = _SUCHE.search(formel)
        if m:
            name = m.group(1).upper().removeprefix("_XLFN.")
            return f"{_DEUTSCH.get(name, name)} findet den Suchwert nicht im Suchbereich"
    return {
        "#REF!": "Bezug zeigt auf eine gelöschte Zeile, Spalte oder ein gelöschtes Blatt",
        "#N/A": "Wert nicht verfügbar",
        "#VALUE!": "falscher Datentyp, oft Text statt Zahl",
        "#DIV/0!": "Teilung durch null oder durch eine leere Zelle",
        "#NAME?": "unbekannter Funktionsname oder Name",
        "#NUM!": "ungültige Zahl",
        "#NULL!": "Bereiche ohne Schnittmenge",
    }.get(fehler, "")


def externe_ziele(wb: Workbook) -> list[str]:
    ziele = []
    for link in getattr(wb, "_external_links", []) or []:
        ziel = getattr(getattr(link, "file_link", None), "Target", None)
        if ziel:
            ziele.append(str(ziel))
    return ziele


def blatt_formeln(ws: Worksheet, werte_ws: Worksheet) -> dict[str, Any]:
    fehler: list[dict[str, Any]] = []
    extern: list[dict[str, Any]] = []
    for row in ws.iter_rows():
        for c in row:
            formel = laden.formeltext(c) if laden.ist_formel(c) else None
            cache = werte_ws[c.coordinate].value if formel else c.value
            ref = c.coordinate
            if isinstance(cache, str) and cache in regeln.FEHLERWERTE:
                fehler.append({
                    "zelle": ref, "fehler": cache, "deutsch": regeln.FEHLERWERTE[cache],
                    "formel": formel, "hinweis": _hinweis(cache, formel),
                })
            elif formel and "#REF!" in formel:
                fehler.append({
                    "zelle": ref, "fehler": "#REF!", "deutsch": "#BEZUG!",
                    "formel": formel, "hinweis": "Formel enthält einen gelöschten Bezug",
                })
            if formel and _EXTERN.search(formel):
                extern.append({"zelle": ref, "formel": formel})

    harte: list[dict[str, Any]] = []
    fehlend: list[dict[str, Any]] = []
    spalten: list[dict[str, Any]] = []
    t = laden.tabelle(ws)
    if t.kopf_zeile is not None:
        datenzeilen = [nr for nr, werte in t.zeilen if not regeln.ist_summenzeile(werte)]
        for off, name in enumerate(t.kopf):
            sp = t.erste_spalte + off
            zellen = [ws.cell(row=nr, column=sp) for nr in datenzeilen]
            gefuellt = [z for z in zellen if laden.ist_formel(z) or not laden.ist_leer(z.value)]
            formeln = [z for z in gefuellt if laden.ist_formel(z)]
            if len(formeln) < regeln.FORMELSPALTE_MIN or len(formeln) / len(gefuellt) < regeln.FORMELSPALTE_ANTEIL:
                continue
            buchstabe = get_column_letter(sp)
            spalten.append({"spalte": buchstabe, "name": name, "formeln": len(formeln), "gefuellt": len(gefuellt)})
            for z in gefuellt:
                if not laden.ist_formel(z):
                    harte.append({"zelle": z.coordinate, "spalte": name, "wert": z.value,
                                  "hinweis": "harter Wert in einer Formelspalte"})
            for z in zellen:
                if laden.ist_formel(z) or not laden.ist_leer(z.value):
                    continue
                fehlend.append({"zelle": z.coordinate, "spalte": name,
                                "hinweis": "Zeile hat Daten, aber die Formel fehlt"})
    return {
        "blatt": ws.title,
        "fehler": fehler,
        "externe_verknuepfungen": extern,
        "formelspalten": spalten,
        "harte_werte": harte,
        "formel_fehlt": fehlend,
    }


def ausfuehren(datei: Path, out: Path) -> dict[str, Any]:
    wb, wv = laden.lade_beide(datei)
    blaetter = [blatt_formeln(ws, wv[ws.title]) for ws in wb.worksheets]
    ziele = externe_ziele(wb)
    daten = {"befehl": "formeln", "datei": datei.name, "externe_dateien": ziele, "blaetter": blaetter}
    zeilen = [f"Datei: `{bericht.zelle_md(datei.name)}`", "",
              "Nur Bericht. Die Datei wird nicht geändert, Formeln werden nicht neu berechnet.", ""]
    if ziele:
        zeilen += ["Externe Dateien: " + ", ".join(f"`{bericht.zelle_md(z)}`" for z in ziele), ""]
    for b in blaetter:
        zeilen += [f"## Blatt {bericht.zelle_md(b['blatt'])}", "", "Fehlerwerte:", ""]
        zeilen += bericht.tabelle_md(["Zelle", "Fehler", "Deutsch", "Formel", "Hinweis"],
                                     [[f["zelle"], f["fehler"], f["deutsch"], f["formel"], f["hinweis"]] for f in b["fehler"]])
        zeilen += ["Harte Werte in Formelspalten:", ""]
        zeilen += bericht.tabelle_md(["Zelle", "Spalte", "Wert"], [[h["zelle"], h["spalte"], h["wert"]] for h in b["harte_werte"]])
        zeilen += ["Formel fehlt in einer Datenzeile:", ""]
        zeilen += bericht.tabelle_md(["Zelle", "Spalte"], [[h["zelle"], h["spalte"]] for h in b["formel_fehlt"]])
        zeilen += ["Externe Verknüpfungen:", ""]
        zeilen += bericht.tabelle_md(["Zelle", "Formel"], [[e["zelle"], e["formel"]] for e in b["externe_verknuepfungen"]])
    json_pfad = bericht.schreibe_bericht(out, f"Formel-Check: {bericht.zelle_md(datei.name)}", zeilen, daten)
    return {
        "befehl": "formeln",
        "bericht": str(out),
        "json": str(json_pfad),
        "fehler": sum(len(b["fehler"]) for b in blaetter),
        "harte_werte": sum(len(b["harte_werte"]) for b in blaetter),
        "formel_fehlt": sum(len(b["formel_fehlt"]) for b in blaetter),
        "externe_verknuepfungen": sum(len(b["externe_verknuepfungen"]) for b in blaetter),
    }
